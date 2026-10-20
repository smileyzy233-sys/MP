from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from copy import deepcopy
import json,hashlib,re
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

B=Path(__file__).resolve().parent
ROOT=Path(r'A:\桌面\study')
TARGET=ROOT/'副窗口讨论交接_20261014.docx'
assert not TARGET.exists(),'A file with the delivery name already exists'
meta=json.loads((B/'source_manifest.json').read_text(encoding='utf-8'))
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0)
sec.left_margin=Cm(2.25);sec.right_margin=Cm(2.25)
sec.header_distance=Cm(0.8);sec.footer_distance=Cm(0.9)
for st in doc.styles:
 if st.type==1:
  pf=st.paragraph_format
  pf.keep_with_next=False;pf.keep_together=False;pf.page_break_before=False
  pf.space_before=Pt(0);pf.space_after=Pt(5);pf.line_spacing=1.18
  st.font.name='Times New Roman';st.font.size=Pt(11)
  st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
for name,size in [('Title',18),('Heading 1',14),('Heading 2',12)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
 st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体')
 st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(7)
for p in sec.header.paragraphs:p.text=''
fp=sec.footer.paragraphs[0];fp.alignment=1
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE')
r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);f.append(r);fp._p.append(f)

def p(text='',style=None):
 q=doc.add_paragraph(text,style)
 q.paragraph_format.keep_with_next=False
 q.paragraph_format.keep_together=False
 q.paragraph_format.page_break_before=False
 return q
def h(text):return p(text,'Heading 1')
def sub(text):return p(text,'Heading 2')
def page():doc.add_page_break()
def table(headers,rows,widths=None):
 tb=doc.add_table(rows=1,cols=len(headers));tb.autofit=False
 for i,s in enumerate(headers):tb.rows[0].cells[i].text=s
 for values in rows:
  cells=tb.add_row().cells
  for i,s in enumerate(values):cells[i].text=s
 if widths:
  for row in tb.rows:
   for cell,w in zip(row.cells,widths):cell.width=Cm(w)
 for ri,row in enumerate(tb.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for cell in row.cells:
   for q in cell.paragraphs:
    q.paragraph_format.space_after=Pt(4);q.paragraph_format.space_before=Pt(4)
    for rr in q.runs:rr.font.size=Pt(10);rr.bold=(ri==0)
  if ri==0:
   for cell in row.cells:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(sh)
 return tb
def mr(txt):
 r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=txt;r.append(t);return r
def ms(base,sub):
 x=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(base));s=OxmlElement('m:sub');s.append(mr(sub));x.append(e);x.append(s);return x
def equation(items):
 q=p();q.alignment=1
 mp=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath')
 for it in items:m.append(mr(it) if isinstance(it,str) else it)
 mp.append(m);q._p.append(mp)
def link(title,url):
 q=p();rel=q.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 a=OxmlElement('w:hyperlink');a.set(qn('r:id'),rel)
 r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');rp.append(c)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp)
 t=OxmlElement('w:t');t.text=title;r.append(t);a.append(r);q._p.append(a)

q=p('企业政策匹配与实证方法\n副窗口讨论交接','Title');q.alignment=1
q=p('2026年10月14日｜供新副窗口继续讨论');q.alignment=1
h('1. 当前进度与资料')
p('最新讨论停在方法选择：用户提出根据企业各年实际行业代码及当时有效政策，构造年度0/1限制状态，用企业面板回归研究其与生产的关系。是否以此替代当前单轮DID尚未决定，成本反演的经验目标也需要随研究问题重新核对。')
p('后续先继续讨论变量含义、识别来源与部门生产口径。本文保存讨论背景和待决问题；文件修改与实证计算按届时明确要求开展。')
table(['资料','用途与本次核对'],[
 ['A:\\桌面\\study\\思路.docx','用户原始想法。已实际读取正文及嵌入图；“实证思路”“反演思路”仍留空。'],
 ['A:\\桌面\\study\\企业微观层面_v3.35.docx','当前方案。本交接生成时重新读取：A.3已补充逐版限制状态、名称匹配与Treat的区别，全文已有GB小类、AMNE部门、模型部门等术语说明。'],
 ['A:\\桌面\\study\\AGENTS.md','项目规则中快照及部分匹配设计仍停留于v3.3。核对实际Word和用户最新想法后使用；本副窗口提出的新方法选择不能由旧设计约定替代。']
],[6.4,10.1])
p('当前Word保留的设定：2011年目录信息公开及随后实施；窗口2007—2014年；2010年为年度动态参照；处理身份按改革前信息固定。此为现有方案状态，用户正在讨论是否改变实证主线。')
p('本副窗口完成了只读文件比较、方法讨论和来源核对，曾给出供主聊天修改的提示词。当前Word内容已与最初检查时不同。本副窗口没有修改研究方案或数据，也没有执行匹配、回归或结构计算。此次新建的交付文件仅为本交接文档。')
sub('文件核对快照')
for item in meta['files'][:2]:
 p(Path(item['path']).name+'\nSHA256：'+item['sha256'])
p('哈希对应本次读取的磁盘保存版本。重新开始讨论时，实际文件如有变化，以新版本内容为准；不要复用较早的版本判断。')

page();h('2. 用户的匹配思路与术语约定')
p('当前可操作的企业匹配资料是税调行业代码。先核实代码所属GB分类版本，取得对应小类名称，再由LLM判断小类名称与政策条款所描述活动是否属于同一类。名称同类判断成立后，将整个小类纳入相应条款覆盖，并赋予属于该小类的企业。真实业务文字是未来可能补充的资料。')
p('政策赋值保持小类名称层级。分类注释中的具体活动、包括与不包括项目，不再用来把同一小类内企业拆分、要求企业补充业务资料或因甲乙活动不同而记为未知。名称关系或分类版本无法确认时，可以保留未知。法律条款原有要求、限定语及生效时间仍须准确记录。')
table(['名称','所指来源或层级'],[
 ['GB行业分类／GB版本','《国民经济行业分类》（GB/T 4754）及其2002、2011、2017等版本。代码版本需依数据来源核实。'],
 ['企业行业代码／GB小类名称','税调当年记录的代码，及对应GB版本正文中的小类名称。门类字母与四位小类数字分别说明。'],
 ['目录条款／条款所述活动','相应版本《外商投资产业指导目录》的条目或子项及其描述的活动。鼓励类、限制类、禁止类是政策分类。'],
 ['ISIC类别／AMNE部门／模型部门','分别指ISIC Rev.4的相应分类层级、所选OECD Analytical AMNE发布版的汇总部门、结构模型部门；三者分别称呼。']
],[5.0,11.5])
p('《思路.docx》嵌入图中的示例是H6361对应“煤炭及制品批发”。它说明LLM输入为小类名称，具体政策命中结论还要有原条款依据。当前Word已加入该名称案例，后续先读取现有A.3，避免重复编造示例结论。')
sub('名称命中、限制状态、实际放宽的区别')
p('名称命中回答小类与条款活动是否同类；限制状态依据条款明确的股比、合资合作、禁止进入等要求记录，鼓励资格另记；实际放宽比较前后具体要求。命中未附准入要求的鼓励类条款，本身不足以记为受限。确认未命中相关限制条款与无法判断匹配关系分别记录。')
p('股比上限提高后仍有限制，可以构成放宽，但受限状态仍可能是1。固定Treat表示本轮实际放宽，年度Restrict表示当期有无限制，两者回答不同问题。')
sub('部门映射路径')
p('企业各期原始代码及版本 → 必要的GB2002—GB2011转换 → GB2017附录B中的GB2011—GB2017对照 → 附录C中的GB2017—ISIC Rev.4对照 → 所选AMNE部门 → 模型部门。已有明确对照直接使用。保留一对多及部分转入：候选落在同一最终部门时可确定部门，跨部门时结合可靠业务或产出依据处理。这一歧义处理属于部门归属路径。')

page();h('3. 固定改革前类别与逐年行业状态')
p('用户举例：企业A在2007—2010年的代码依次为A1001、A1001、G1003、H7422，在2007版目录下均受限；2011—2014年均为H7422，在2011版目录下限制放松。示例代码用于讨论逻辑，未核验为真实企业记录。')
p('在这个例子中，两种思路可以得到同一处理结论：以2010年的H7422类别为共同对象，比较两版目录的要求，如果该类别实际放宽，则固定Treat为1。更早年份的其他代码仍可用于年度状态描述。相同代码跨GB版本的含义也需要核实。')
p('区别出现在企业改革后转行业：若2010年属于一直受限的X类，2011年转入一直开放的Y类，两版政策对X和Y均未改变，则企业年度状态会从1变0，变化来源却是行业选择。固定改革前类别方法与按实际当期类别变化的方法，在这种情况下有不同研究含义。')
sub('用户最新提出的面板方案')
p('企业各年当期代码 → 核实GB版本 → 取得当期GB小类名称 → 与当时相应政策文件匹配 → 有限制记1、确认无限制记0 → 用年度限制状态作为面板回归解释变量。用户询问此方案是否更合适，尚未决定替代当前DID。')
p('0/1自变量可以进入企业面板回归，生产金额也可以保留PPML。DID是识别策略，面板是数据结构；变更变量和比较对象，比变更回归名称更重要。一个示意线性规格为：')
equation([ms('Y','ft'),' = ',ms('α','f'),' + ',ms('λ','t'),' + β',ms('Restrict','ft'),' + ',ms('X','ft'),'′γ + ',ms('ε','ft')])
p('企业固定效应处理企业稳定差异，年份固定效应处理共同年度变化。因果解释需要说明限制状态变化与未观察生产冲击的关系。企业根据经营表现转行业，或政策选择性开放某些行业，都可能影响这一解释。')
p('在“行业固定、只有一轮共同时间的1→0政策变化、其他状态不变”的特殊情况下，有如下变量关系：')
equation([ms('Restrict','ft'),' = ',ms('Restrict','f,pre'),' − ',ms('Treat','f'),' × ',ms('Post','t')])
p('改革前状态被企业固定效应吸收后，在相同样本、控制与共同系数设定下，状态回归与相应DID是不同写法，系数符号相反。这个等价关系以所列条件为前提。')
sub('状态变化应区分四种来源')
table(['来源','讨论中的含义'],[
 ['真实政策取消限制，1→0','包含可用于政策识别的变化。'],
 ['企业转行业导致1→0或0→1','包含企业选择，需区别于政策变化。'],
 ['GB分类修订或匹配调整','可能改变观察到的标签，需核查真实变化。'],
 ['要求放宽但仍有其他限制，1→1','简单0/1状态不能记录该次放宽。']
],[7.5,9.0])

page();h('4. 有效比较、时间、样本与反演')
sub('固定效应与政策时间')
p('状态全期不变的企业，其稳定差异会被企业固定效应吸收；所有企业同年同幅变化会被年份固定效应吸收。若当期同一GB小类、年份内共用一个限制状态，当期GB小类×年份固定效应会完全吸收该变量。需要检查固定效应后剩余的状态变化，并按政策赋值及误差相关结构确定聚类。')
p('年度法律状态需按实际有效规则定义。2007版目录自2007年12月1日施行，此前适用规则需要补核；2011版目录自2012年1月30日施行。2011年信息公开反应与法律实施状态分别记录。年初、年末状态或转换年处理如何与全年生产对应，尚待选择。来源见第5页。')
sub('“稳定企业”与选择偏差')
p('DID与一般面板回归均不普遍要求企业2007—2014年八期齐全。企业标识稳定、改革前已经经营、全期被调查到和行业全期不变，是不同概念。按政策后的存续或全期出现筛选，可能选择了受政策影响的企业构成；平衡样本不会自动保证因果识别，非平衡样本也不会自动解决失访。')
p('企业在税调消失，可能是停产、退出调查、漏报或标识变化。确认该期未生产时，生产价值可以记零；年内退出仍可能有当年产出。调查缺失保留缺失；并购、更名或标识变化先核实记录延续，避免把主体变化当成生产消失。')
p('当前Word A.2.3已允许非连续观察，研究改革前已有企业。后续应分别核查估计样本、失访原因和总体覆盖；完整观测样本可作为补充比较，并说明其对应对象。')
sub('企业效应与部门总生产的衔接')
p('某部门2014年总生产包括改革前已有企业当期生产，以及改革后新设企业当期生产；两部分都还可能含税调未覆盖企业。新进入、退出、控制权转换和跨部门迁移会改变部门结构。存量样本的平均相对效应不自动等于完整部门生产变化。')
p('反演需要同一部门、控制定义和价格口径下“有改革”和“无本轮改革”的生产。补充实际行业或AMNE资料能帮助恢复实际总体，但新进入及未覆盖生产的无改革值仍需要识别依据。把它们的实际值同时用于无改革情景，等于加入这些生产不受改革影响的假设。当前Word A.6.3已记录这一问题。')
p('以年度限制状态关联回归预测“把Restrict设为0”的生产，需要核对因果解释、比较支持、外推范围和企业行业选择。用关联目标反演的成本属于给定回归和模型假设的情景解释；其与2011年改革造成的成本变化是否一致，需要另外论证。')

page();h('5. 下次讨论从哪里继续')
p('最近一次建议是先建设可复用的年度类别政策暴露底表，再按研究问题与有效变化选择实证规格。用户尚未接受该建议，也未决定放弃DID、改变年份、采用新估计器或调整成本框架。')
sub('优先待决问题')
table(['问题','需要形成的判断'],[
 ['主要研究对象','继续识别2011年实际放宽，还是转向年度有无限制与企业经营表现的关系？关联研究与因果研究分别说明。'],
 ['核心解释变量','以年度Restrict、实际要求变化或固定Treat为主？如何覆盖仍受限情况下的放宽？'],
 ['有效识别来源','政策变化、企业转行业、分类修订及匹配调整分别贡献哪些观察变化？固定效应后剩余信息是什么？'],
 ['年度规则与样本','如何处理公开与实施、转换年、失访和新设企业？哪些信息还需从税调来源核实？'],
 ['反演对象','如何形成完整或限定覆盖范围的部门反事实？对未覆盖生产、进入和行业选择增加什么假设？']
],[5.0,11.5])
p('目前尚未统计：行业版本全样本覆盖、年度0/1转换数量、转换来源、平衡与非平衡样本规模、退出和失访比例、部门覆盖率、政策组数。上述均是后续可核查事项，本次没有新估计或数值结果。')
sub('可以发给新副窗口的起始消息')
p('请先读取这份交接文档，并实际核对《思路.docx》和当前《企业微观层面_v3.35.docx》的有关内容。继续讨论我提出的“按企业各年实际行业代码和当时政策构造0/1限制状态面板”的可行性。实证和反演仍在形成中。请比较年度状态回归与固定改革前类别DID的研究对象、识别来源、行业选择、样本覆盖和部门反事实，依据我接下来的问题推进。现阶段先讨论方法，不修改研究方案或执行数据与模型计算。')
sub('已查看的方法与政策来源')
link('Callaway与Sant’Anna：did方法说明（非平衡面板及缺失原因）','https://bcallaway11.github.io/did/articles/did-basics.html')
link('Rathnayake等：Difference-in-Differences with Sample Selection','https://arxiv.org/abs/2411.09221')
link('Imai与Kim（2019）：固定效应因果解释与纵向数据','https://imai.fas.harvard.edu/research/files/FEmatch.pdf')
link('国家发展改革委：2007版目录第57号令及施行日期','https://zfxxgk.ndrc.gov.cn/web/iteminfo.jsp?id=603')
link('国家发展改革委：2011版目录第12号令及施行日期','https://www.ndrc.gov.cn/xxgk/zcfb/fzggwl/201112/t20111229_960737.html')
p('技能与核验：本交接使用research-handoff-memory保存来源、状态与待决问题，使用docx生成Word。按用户要求交付一份独立Word，不替换主聊天的交接包或项目规则。原生公式、空页眉、页码及分页设置另作成品核验。')

doc.core_properties.title='企业政策匹配与实证方法：副窗口讨论交接'
doc.core_properties.subject='年度限制状态面板、单轮DID与部门成本反演的待决讨论'
doc.save(TARGET)
# Disable pagination controls throughout all XML parts, including inherited styles.
with ZipFile(TARGET) as z:blobs={n:z.read(n) for n in z.namelist()}
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
for name,data in list(blobs.items()):
 if name.startswith('word/') and name.endswith('.xml'):
  tree=etree.fromstring(data)
  for tag in ['keepNext','keepLines','pageBreakBefore']:
   for e in tree.findall('.//w:'+tag,ns):e.set(qn('w:val'),'0')
  blobs[name]=etree.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
with ZipFile(TARGET,'w',ZIP_DEFLATED) as z:
 for n,data in blobs.items():z.writestr(n,data)
for item in meta['files']:
 assert hashlib.sha256(Path(item['path']).read_bytes()).hexdigest()==item['sha256'],'Protected source changed during build'
print(json.dumps({'target':str(TARGET),'sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),'size':TARGET.stat().st_size},ensure_ascii=False))
