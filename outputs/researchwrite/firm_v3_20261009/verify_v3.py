from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
import hashlib, json, re
import pdfplumber

ROOT = Path(__file__).resolve().parent
TARGET = ROOT.parents[2] / '企业微观层面_v3.docx'
RENDER = ROOT / 'render/round4'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
RNS = '{http://schemas.openxmlformats.org/package/2006/relationships}'
checks = []

def check(name, ok, detail=None):
    checks.append({'check': name, 'passed': bool(ok), 'detail': detail})

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

meta = json.loads((ROOT / 'build_meta.json').read_text(encoding='utf-8'))
manifest = json.loads((ROOT / 'source_manifest.json').read_text(encoding='utf-8'))
for item in manifest:
    check('source_unchanged_' + item['label'], sha(item['path']) == item['sha256'], item['path'])
check('output_matches_build', sha(TARGET) == meta['sha256'], sha(TARGET))
check('draft_matches_build', sha(ROOT / 'proposal_v3.md') == meta['draft_sha256'])

with ZipFile(TARGET) as z:
    check('zip_integrity', z.testzip() is None)
    roots = {n: etree.fromstring(z.read(n)) for n in z.namelist() if n.endswith(('.xml', '.rels'))}
    main = roots['word/document.xml']
    notes = roots['word/footnotes.xml']
    fulltext = ''.join(main.xpath('//w:t/text()|//m:t/text()', namespaces=NS))
    check('no_raster_or_legacy_objects', not main.xpath('//w:drawing|//w:pict', namespaces=NS))
    check('native_math_count', len(main.xpath('//m:oMath', namespaces=NS)) == 116)
    math_records = json.loads((ROOT / 'equations.json').read_text(encoding='utf-8'))
    actual_math = [''.join(x.xpath('.//m:t/text()', namespaces=NS)) for x in main.xpath('//m:oMath', namespaces=NS)]
    check('math_text_matches_builder', actual_math == [x['text'] for x in math_records])
    check('native_math_structure', all(main.xpath('//m:' + tag, namespaces=NS) for tag in ['f', 'acc', 'sSub', 'sSup', 'nary']))
    bad = ['activity_id', 'version_link', '[[FLOWCHART]]', 'PLACEHOLDER', '\\tag', '$', '\ufffd']
    check('no_placeholders_or_crossversion_ids', not any(s in fulltext for s in bad))
    check('six_dimensions_preserved', all(s in fulltext for s in ['治理限制', '其他有限限制', '地域或单一外商持股']))
    display_labels = []
    for p in main.xpath('/w:document/w:body/w:p', namespaces=NS):
        if p.xpath('./w:pPr/w:pStyle[@w:val="Equation"]', namespaces=NS) and p.xpath('.//m:oMath', namespaces=NS):
            txt = ''.join(p.xpath('.//w:t/text()', namespaces=NS))
            display_labels.extend(re.findall(r'\((A\d+)\)', txt))
    check('display_equations_A1_to_A17', display_labels == ['A' + str(i) for i in range(1, 18)], display_labels)
    refs = re.findall(r'（(A\d+)）', fulltext)
    check('equation_references_resolve', set(refs).issubset(set(display_labels)), sorted(set(refs)))
    check('caption_counts', all(fulltext.count(x) == 1 for x in meta['captions']) and fulltext.count('图A1 企业估计与部门成本反演的衔接') == 1)
    ref_ids = main.xpath('//w:footnoteReference/@w:id', namespaces=NS)
    note_ids = notes.xpath('//w:footnote[not(@w:type)]/@w:id', namespaces=NS)
    check('footnote_references', ref_ids == ['1', '2', '3'] and sorted(note_ids) == ['1', '2', '3'])
    for part in ['document', 'footnotes']:
        rel = roots['word/_rels/' + part + '.xml.rels']
        mapping = {x.get('Id'): x for x in rel}
        hyper_refs = roots['word/' + part + '.xml'].xpath('//w:hyperlink/@r:id', namespaces=NS)
        check(part + '_hyperlinks', all(r in mapping and mapping[r].get('TargetMode') == 'External' and mapping[r].get('Target', '').startswith('https://') for r in hyper_refs), len(hyper_refs))

doc = Document(TARGET)
heads = [p.text for p in doc.paragraphs if p.style.name in ['Heading 1', 'Heading 2']]
check('heading_sequence', heads == meta['headings'])
check('seven_main_sections', len([s for s in heads if re.match(r'^A\.[1-7] ', s)]) == 7)
check('sixteen_subsections', len([s for s in heads if re.match(r'^A\.\d\.\d ', s)]) == 16)
section_ids = {s.split(' ', 1)[0] for s in heads if s.startswith('A.')}
check('section_references_resolve', set(re.findall(r'A\.\d+(?:\.\d+)?', fulltext)).issubset(section_ids))
check('four_editable_tables', len(doc.tables) == 4)
sections = [tuple(int(getattr(s, k)) for k in ['page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin']) for s in doc.sections]
src = Document(manifest[0]['path'])
original_sections = [tuple(int(getattr(s, k)) for k in ['page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin']) for s in src.sections]
check('source_page_geometry_retained', sections == original_sections)
render_meta = json.loads((RENDER / 'render_meta.json').read_text(encoding='utf-8-sig'))
check('opened_by_Microsoft_Word', render_meta['input'] == str(TARGET) and render_meta['math'] == 116 and render_meta['footnotes'] == 3 and render_meta['tables'] == 4)
with pdfplumber.open(RENDER / 'document.pdf') as pdf:
    check('fourteen_rendered_pages', len(pdf.pages) == 14 and render_meta['pages'] == 14)
    rendered = [p.extract_text() or '' for p in pdf.pages]
    check('all_pages_searchable', all(len(t.strip()) > 100 for t in rendered))
    check('no_replacement_glyphs', all('\ufffd' not in t for t in rendered))
    check('all_numbered_equations_rendered', all(re.search(r'\(\s*A' + str(i) + r'\s*\)', '\n'.join(rendered)) for i in range(1,18)))
    check('footnote_3_not_isolated', '\n3\n' not in rendered[6])
check('final_render_after_build', (RENDER / 'document.pdf').stat().st_mtime >= TARGET.stat().st_mtime)
result = {'output': str(TARGET), 'sha256': sha(TARGET), 'render': str(RENDER),
          'pdf_sha256': sha(RENDER / 'document.pdf'), 'all_passed': all(c['passed'] for c in checks), 'checks': checks}
(ROOT / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'all_passed': result['all_passed'], 'checks': len(checks), 'failed': [x for x in checks if not x['passed']]}, ensure_ascii=False))
raise SystemExit(0 if result['all_passed'] else 1)
