from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
import re,json,hashlib,importlib.util,copy
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
PROJECT=ROOT.parents[2]
SOURCE=PROJECT/'企业微观层面_新方案_框架重写版.docx'
TARGET=PROJECT/'企业微观层面_v3.docx'
helper_path=PROJECT/'outputs/researchwrite/firm_old_framework_20261007/build_document.py'
spec=importlib.util.spec_from_file_location('v3_math_helper',helper_path)
helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
helper.helper.GREEK.update({'chi':'χ','ell':'ℓ'})
NS=helper.NS
RNS='http://schemas.openxmlformats.org/package/2006/relationships'
WNS=NS['w']
records=[]; foot_refs={}; fn_links={}; foot_defs={}
DOC=None

def native_math(tex,display=False):
    obj=helper.native_math(tex,display)
    if not display:
        box=OxmlElement('m:box');pr=OxmlElement('m:boxPr');nb=OxmlElement('m:noBreak');nb.set(qn('m:val'),'1');pr.append(nb);box.append(pr);content=OxmlElement('m:e')
        for child in list(obj):obj.remove(child);content.append(child)
        box.append(content);obj.append(box)
    records.append({'latex':tex,'display':display,'text':''.join(obj.xpath('.//m:t/text()',namespaces=NS))})
    return obj

def append_hyperlink(p,label,url,foot=False):
    h=OxmlElement('w:hyperlink')
    if foot:
        key=f'rId{len(fn_links)+1}';fn_links[key]=url
    else:key=DOC.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h.set(qn('r:id'),key)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    c=OxmlElement('w:color');c.set(qn('w:val'),'000000');pr.append(c)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p,text,foot=False):
    pattern=r'(\$[^$]+\$|`[^`]+`|\[\^[^\]]+\]|\[[^\]]+\]\(https?://[^)]+\))'
    for part in re.split(pattern,text):
        if not part:continue
        if part.startswith('$') and part.endswith('$'):
            p._p.append(native_math(part[1:-1]));continue
        if part.startswith('[^'):
            label=part[2:-1]
            if label not in foot_refs:foot_refs[label]=len(foot_refs)+1
            p.add_run('\u2060')
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');st=OxmlElement('w:rStyle');st.set(qn('w:val'),'FootnoteReference');pr.append(st);r.append(pr)
            ref=OxmlElement('w:footnoteReference');ref.set(qn('w:id'),str(foot_refs[label]));r.append(ref);p._p.append(r);continue
        if part.startswith('[') and '](' in part:
            label,url=re.fullmatch(r'\[([^\]]+)\]\((.+)\)',part).groups();append_hyperlink(p,label,url,foot);continue
        code=part.startswith('`') and part.endswith('`')
        run=p.add_run(part[1:-1] if code else part)
        if code:run.font.name='Consolas';run.font.size=Pt(10.5)

def borders(cell,color='8497AB',fill='FFFFFF'):
    pr=cell._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
    for s in ['top','left','bottom','right']:
        e=OxmlElement('w:'+s);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),color);bd.append(e)
    pr.append(bd);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);pr.append(sh)

def flowchart(doc):
    tab=doc.add_table(rows=8,cols=2);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    for col in tab.columns:col.width=Inches(3.25)
    for row in tab.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.width=Inches(3.25);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side,val in [('top',65),('bottom',65),('left',95),('right',95)]:
                n=OxmlElement('w:'+side);n.set(qn('w:w'),str(val));n.set(qn('w:type'),'dxa');m.append(n)
            pr.append(m)
    rows={0:'政策逐版编码与GB2017行业匹配\n大语言模型辅助识别，人工核查法律适用范围',1:'↓',2:'固定企业改革前业务，逐版判断政策待遇\n确定改革事件及适用范围',3:'↓',6:'核对生产覆盖、部门归属及模型比较范围\n有完整生产依据或明确补充假设时，形成部门目标',7:'↓\n部门MP成本反演 → 一般均衡反事实'}
    for idx,text in rows.items():
        cell=tab.cell(idx,0).merge(tab.cell(idx,1));p=cell.paragraphs[0];p.text=text
        if idx not in [1,3]:borders(cell,fill='EEF2F6' if idx in [0,2,6] else 'FFFFFF')
    tab.cell(4,0).text='企业—年份基准分析\n原有企业的规模、控制权与存续'
    tab.cell(4,1).text='GB四位小类—地区—年份补充分析\n包含新设企业的完整行业生产'
    tab.cell(5,0).text='解释原有企业的调整\n核实原有企业产出在部门中的覆盖范围'
    tab.cell(5,1).text='分别估计外国和国内控制生产\n按真实产出加总至模型部门'
    for ri in [4,5]:
        for cell in tab.rows[ri].cells:borders(cell)
    seen=set()
    for row in tab.rows:
        for cell in row.cells:
            if cell._tc in seen:continue
            seen.add(cell._tc)
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
                p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.first_line_indent=Pt(0)
                for run in p.runs:run.font.size=Pt(10.5)
    p=doc.add_paragraph('图A1 企业估计与部门成本反演的衔接',style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=False
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)

def add_footnotes(path):
    from docx.text.paragraph import Paragraph
    root=etree.Element(qn('w:footnotes'),nsmap={'w':WNS,'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'})
    for fid,kind in [(-1,'separator'),(0,'continuationSeparator')]:
        fn=etree.SubElement(root,qn('w:footnote'));fn.set(qn('w:type'),kind);fn.set(qn('w:id'),str(fid))
        p=etree.SubElement(fn,qn('w:p'));r=etree.SubElement(p,qn('w:r'));etree.SubElement(r,qn('w:'+kind))
    for label,fid in foot_refs.items():
        fn=etree.SubElement(root,qn('w:footnote'));fn.set(qn('w:id'),str(fid));p=OxmlElement('w:p');fn.append(p)
        pp=OxmlElement('w:pPr');st=OxmlElement('w:pStyle');st.set(qn('w:val'),'FootnoteText');pp.append(st);p.append(pp)
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sty=OxmlElement('w:rStyle');sty.set(qn('w:val'),'FootnoteReference');pr.append(sty);r.append(pr);r.append(OxmlElement('w:footnoteRef'));p.append(r)
        proxy=Paragraph(p,DOC._body);inline(proxy,' '+foot_defs[label],foot=True)
    with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    parts['word/footnotes.xml']=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    relroot=etree.Element('{'+RNS+'}Relationships',nsmap={None:RNS})
    for key,url in fn_links.items():etree.SubElement(relroot,'{'+RNS+'}Relationship',Id=key,Type=RT.HYPERLINK,Target=url,TargetMode='External')
    parts['word/_rels/footnotes.xml.rels']=etree.tostring(relroot,xml_declaration=True,encoding='UTF-8',standalone=True)
    rels=etree.fromstring(parts['word/_rels/document.xml.rels'])
    if not any(x.get('Type')==RT.FOOTNOTES for x in rels):
        ids={x.get('Id') for x in rels};idx=1
        while f'rId{idx}' in ids:idx+=1
        etree.SubElement(rels,'{'+RNS+'}Relationship',Id=f'rId{idx}',Type=RT.FOOTNOTES,Target='footnotes.xml')
    parts['word/_rels/document.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
    ct=etree.fromstring(parts['[Content_Types].xml']);ctns=etree.QName(ct).namespace
    if not any(x.get('PartName')=='/word/footnotes.xml' for x in ct):etree.SubElement(ct,'{'+ctns+'}Override',PartName='/word/footnotes.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)

def build():
    global DOC,foot_defs
    manifest=json.loads((ROOT/'source_manifest.json').read_text(encoding='utf-8'))
    for item in manifest:assert hashlib.sha256(Path(item['path']).read_bytes()).hexdigest()==item['sha256'],item['path']+' changed'
    if TARGET.exists():
        meta_path=ROOT/'build_meta.json';assert meta_path.exists(),'Output already exists'
        meta=json.loads(meta_path.read_text(encoding='utf-8'));assert hashlib.sha256(TARGET.read_bytes()).hexdigest()==meta['sha256'],'Output changed externally'
    draft=(ROOT/'proposal_v3.md').read_text(encoding='utf-8')
    foot_defs=dict(re.findall(r'^\[\^([^\]]+)\]: (.+)$',draft,re.M));draft=re.sub(r'^\[\^[^\]]+\]: .+$','',draft,flags=re.M)
    DOC=Document(SOURCE);body=DOC.element.body
    for child in list(body):
        if etree.QName(child).localname!='sectPr':body.remove(child)
    DOC.core_properties.title='附录A 外资准入负面清单的政策测度、跨国投资成本识别与一般均衡反事实'
    DOC.core_properties.author='';DOC.core_properties.last_modified_by='';DOC.core_properties.comments=''
    eq=DOC.styles['Equation'].paragraph_format;eq.tab_stops.clear_all();eq.tab_stops.add_tab_stop(Inches(3.02),WD_TAB_ALIGNMENT.CENTER);eq.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    for name in ['Title','Heading 1','Heading 2']:
        DOC.styles[name].font.color.rgb=RGBColor(0,0,0)
    if 'Footnote Text' not in DOC.styles:DOC.styles.add_style('Footnote Text',1)
    ft=DOC.styles['Footnote Text'];helper.helper.font(ft,9.5);ft.paragraph_format.line_spacing=1.0;ft.paragraph_format.space_after=Pt(1);ft.paragraph_format.first_line_indent=Pt(0);ft.paragraph_format.keep_together=True
    if 'Footnote Reference' not in DOC.styles:DOC.styles.add_style('Footnote Reference',2)
    DOC.styles['Footnote Reference'].font.superscript=True
    lines=draft.splitlines();i=0;refs=False;labels=[];heads=[];captions=[];table_count=0
    helper.helper.inline=inline
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line=='[[FLOWCHART]]':flowchart(DOC);continue
        if line.startswith('# '):
            p=DOC.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;inline(p,line[2:]);continue
        if line.startswith('## '):
            h=line[3:];refs=h.startswith('参考文献');DOC.add_heading(h,1);heads.append(h);continue
        if line.startswith('### '):
            h=line[4:];DOC.add_heading(h,2);heads.append(h);continue
        if line.startswith('$$'):
            expr=line[2:]
            while '$$' not in expr:expr+='\n'+lines[i];i+=1
            expr=expr.rsplit('$$',1)[0].strip();m=re.search(r'\\tag\{([^}]+)\}',expr);assert m;label=m[1];expr=re.sub(r'\\tag\{[^}]+\}','',expr).strip()
            prev=DOC.paragraphs[-1]
            if prev.text.rstrip().endswith(('：',':')):prev.paragraph_format.keep_with_next=True
            p=DOC.add_paragraph(style='Equation');p.paragraph_format.keep_with_next=True;p.add_run('\t');p._p.append(native_math(expr,True));p.add_run('\t('+label+')');labels.append(label);continue
        if line.startswith('|'):
            rows=[]
            while True:
                cells=[x.strip() for x in line.strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x or '') for x in cells):rows.append(cells)
                if i>=len(lines) or not lines[i].strip().startswith('|'):break
                line=lines[i].strip();i+=1
            helper.helper.add_table(DOC,rows);table_count+=1
            table=DOC.tables[-1]
            for tag in ['top','bottom']:
                el=table._tbl.tblPr.find(qn('w:tblCellMar')).find(qn('w:'+tag));el.set(qn('w:w'),'50')
            ratios=([.26,.74] if table_count==1 else [.12,.27,.61] if table_count==2 else [.19,.38,.43])
            for col,ratio in zip(table.columns,ratios):col.width=int(Inches(6.5)*ratio)
            for ri,row in enumerate(table.rows):
                for ci,cell in enumerate(row.cells):
                    cell.width=int(Inches(6.5)*ratios[ci])
                    if table_count==2:
                        for para in cell.paragraphs:para.paragraph_format.keep_with_next=(ri<len(table.rows)-1)
            continue
        if re.match(r'^表A\d+',line):
            p=DOC.add_paragraph(style='Caption');inline(p,line);p.paragraph_format.keep_with_next=True;captions.append(line);continue
        p=DOC.add_paragraph(style='Reference' if refs else 'Normal');inline(p,line)
    assert labels==[f'A{i}' for i in range(1,18)],labels
    assert table_count==3
    pending=ROOT/'v3_pending.docx';DOC.save(pending);add_footnotes(pending)
    with ZipFile(pending) as z:
        assert z.testzip() is None
        xml=etree.fromstring(z.read('word/document.xml'));math_count=len(xml.xpath('//m:oMath',namespaces=NS));assert math_count==len(records)
        assert not xml.xpath('//w:drawing|//w:pict',namespaces=NS),'No raster equations or drawings intended'
        assert len(xml.xpath('//w:footnoteReference',namespaces=NS))==3
    pending.replace(TARGET)
    meta={'output':str(TARGET),'sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),'bytes':TARGET.stat().st_size,'equations':labels,'native_math_objects':len(records),'substantive_tables':table_count,'editable_flowcharts':1,'footnotes':len(foot_refs),'headings':heads,'captions':captions,'draft_sha256':hashlib.sha256((ROOT/'proposal_v3.md').read_bytes()).hexdigest()}
    (ROOT/'build_meta.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8');(ROOT/'equations.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(meta,ensure_ascii=False,indent=2))

if __name__=='__main__':build()
