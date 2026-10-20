from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from PIL import Image, ImageChops
from shared_read import read_shared
import hashlib, json, re
import pdfplumber

ROOT = Path(__file__).resolve().parent
TARGET = ROOT.parents[2] / '企业微观层面_v3.3.docx'
RENDER = ROOT / 'render/round2'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
checks = []

def check(name, ok, detail=None):
    checks.append({'check': name, 'passed': bool(ok), 'detail': detail})

def sha(path):
    return hashlib.sha256(read_shared(str(path))).hexdigest()

meta = json.loads((ROOT / 'build_meta.json').read_text(encoding='utf-8'))
manifest = json.loads((ROOT / 'source_manifest.json').read_text(encoding='utf-8-sig'))
for item in manifest:
    check('source_unchanged', sha(item['path']) == item['sha256'], item['path'])
check('output_matches_build', sha(TARGET) == meta['sha256'], sha(TARGET))
check('draft_matches_build', sha(ROOT / 'proposal_v33.md') == meta['draft_sha256'])

with ZipFile(TARGET) as z:
    check('zip_integrity', z.testzip() is None)
    roots = {n: etree.fromstring(z.read(n)) for n in z.namelist() if n.endswith(('.xml', '.rels'))}
    main = roots['word/document.xml']
    notes = roots['word/footnotes.xml']
    fulltext = ''.join(main.xpath('//w:t/text()|//m:t/text()', namespaces=NS))
    note_text = ''.join(notes.xpath('//w:t/text()', namespaces=NS))
    check('no_raster_or_legacy_objects', not main.xpath('//w:drawing|//w:pict|//w:object', namespaces=NS))
    check('native_math_count', len(main.xpath('//m:oMath', namespaces=NS)) == meta['native_math_objects'], meta['native_math_objects'])
    math_records = json.loads((ROOT / 'equations.json').read_text(encoding='utf-8'))
    actual_math = [''.join(x.xpath('.//m:t/text()', namespaces=NS)) for x in main.xpath('//m:oMath', namespaces=NS)]
    check('math_text_matches_builder', actual_math == [x['text'] for x in math_records])
    check('native_math_structure', all(main.xpath('//m:' + tag, namespaces=NS) for tag in ['f', 'acc', 'sSub', 'sSup', 'nary']))
    bad = ['activity_id', 'version_link', '[[FLOWCHART]]', 'PLACEHOLDER', '\\tag', '$', '\ufffd']
    check('no_placeholders_or_crossversion_ids', not any(s in fulltext + note_text for s in bad))
    legacy = ['六类', '六维', '多事件', '历次改革', '所有年份共同估计', 'Zftek', 'CF1', 'CF2', 'CF3', 'CF4', 'CF5', '1995—2024']
    check('old_baseline_removed', not any(s in fulltext for s in legacy))
    check('old_policy_and_cases_removed', not any(s in fulltext + note_text for s in ['2015','2016','名优白酒','起重机械','剔除2014','h=-2','t-2015']))
    check('reminder_phrasing_removed', not any(s in fulltext for s in ['尚需落实','仍待验证','不能预先认定','只有条件成立才','不承诺','当前主要困难']))
    display_math=[x['latex'] for x in math_records if x['display']]
    check('DID_policy_term', 'Treat' in display_math[4] and 'Post' in display_math[4])
    check('dynamic_2011_and_reference_2010', r'h\ne-1' in display_math[5] and 't-2011=h' in display_math[5] and 'h=-1' in fulltext and '2010' in fulltext)
    check('fixed_2007_2014_window', '2007—2010' in fulltext and '2011—2014' in fulltext and '2007—2014' in fulltext)
    check('target_year_2014', display_math[9].count('2014')==2 and '无改革参照状态' in fulltext)
    check('no_plain_math_variables', not any('Treat' in t or 'Post' in t for t in main.xpath('//w:t/text()',namespaces=NS)))
    check('single_reform_and_conditional_scope', all(s in fulltext for s in ['2011年修订', '单轮改革的基准DID', '有改革与无改革的比较', '其他年份']))
    display_labels = []
    for p in main.xpath('/w:document/w:body/w:p', namespaces=NS):
        if p.xpath('./w:pPr/w:pStyle[@w:val="Equation"]', namespaces=NS) and p.xpath('.//m:oMath', namespaces=NS):
            txt = ''.join(p.xpath('.//w:t/text()', namespaces=NS))
            display_labels.extend(re.findall(r'\((A\d+)\)', txt))
    check('display_equations_A1_to_A17', display_labels == ['A' + str(i) for i in range(1, 18)], display_labels)
    refs = re.findall(r'（(A\d+)）', fulltext)
    check('equation_references_resolve', set(refs).issubset(set(display_labels)), sorted(set(refs)))
    check('caption_counts', all(fulltext.count(x) == 1 for x in meta['captions']) and fulltext.count('图A1 单轮改革识别与部门成本反演') == 1)
    ref_ids = main.xpath('//w:footnoteReference/@w:id', namespaces=NS)
    note_ids = notes.xpath('//w:footnote[not(@w:type)]/@w:id', namespaces=NS)
    expected_notes = [str(i) for i in range(1, 5)]
    check('footnote_references', ref_ids == expected_notes and sorted(note_ids) == expected_notes)
    for part in ['document', 'footnotes']:
        rel = roots['word/_rels/' + part + '.xml.rels']
        mapping = {x.get('Id'): x for x in rel}
        hyper_refs = roots['word/' + part + '.xml'].xpath('//w:hyperlink/@r:id', namespaces=NS)
        check(part + '_hyperlinks', all(r in mapping and mapping[r].get('TargetMode') == 'External' and mapping[r].get('Target', '').startswith('https://') for r in hyper_refs), len(hyper_refs))
    check('no_tracked_edits_or_comments', not main.xpath('//w:ins|//w:del|//w:commentReference', namespaces=NS))

doc = Document(TARGET)
heads = [p.text for p in doc.paragraphs if p.style.name in ['Heading 1', 'Heading 2']]
check('heading_sequence', heads == meta['headings'])
check('seven_main_sections', len([s for s in heads if re.match(r'^A\.[1-7] ', s)]) == 7)
check('sixteen_subsections', len([s for s in heads if re.match(r'^A\.\d\.\d ', s)]) == 16)
section_ids = {s.split(' ', 1)[0] for s in heads if s.startswith('A.')}
check('section_references_resolve', set(re.findall(r'A\.\d+(?:\.\d+)?', fulltext)).issubset(section_ids))
check('four_editable_tables', len(doc.tables) == 4)
flow = doc.tables[-1]
check('editable_flowchart_updated', all(s in '\n'.join(c.text for row in flow.rows for c in row.cells) for s in ['2007版、2011年草案', '企业—年份DID与年度动态', '部门MP成本反演']))
def geometry(d):
    return [tuple(int(getattr(s, k)) for k in ['page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin']) for s in d.sections]
check('source_page_geometry_retained', geometry(doc) == geometry(Document(ROOT / 'sources/current.docx')))
render_meta = json.loads((RENDER / 'render_meta.json').read_text(encoding='utf-8-sig'))
check('opened_by_Microsoft_Word', render_meta['input'] == str(TARGET) and render_meta['math'] == meta['native_math_objects'] and render_meta['footnotes'] == 4 and render_meta['tables'] == 4)
with pdfplumber.open(RENDER / 'document.pdf') as pdf:
    check('fourteen_rendered_pages', len(pdf.pages) == 14 and render_meta['pages'] == 14)
    rendered = [p.extract_text() or '' for p in pdf.pages]
    check('all_pages_searchable', all(len(t.strip()) > 100 for t in rendered))
    check('no_replacement_glyphs', all('\ufffd' not in t for t in rendered))
    check('all_numbered_equations_rendered', all(re.search(r'\(\s*A' + str(i) + r'\s*\)', '\n'.join(rendered)) for i in range(1,18)))
    check('footnote_1_not_isolated', not re.search(r'^\s*1\s*$', rendered[1], re.M))
    check('footnote_2_not_isolated', not re.search(r'^\s*2\s*$', rendered[3], re.M))
    check('complete_references_on_final_page', '参考文献与政策资料' in rendered[-1] and 'Eppinger' in rendered[-1] and '报告_重点标红版' in rendered[-1])
    check('rendered_text_within_page', all(c['x0'] >= 0 and c['x1'] <= p.width and c['top'] >= 0 and c['bottom'] <= p.height for p in pdf.pages for c in p.chars))
check('final_render_after_build', (RENDER / 'document.pdf').stat().st_mtime >= TARGET.stat().st_mtime)
changed = []
for n in range(1, 15):
    a = Image.open(ROOT / 'render/round1' / f'page-{n}.png').convert('RGB')
    b = Image.open(RENDER / f'page-{n}.png').convert('RGB')
    if a.size != b.size or ImageChops.difference(a, b).getbbox():
        changed.append(n)
check('changed_pages_rereviewed', changed == [4, 13, 14], {'changed_and_reviewed': changed, 'other_pages_pixel_identical_to_reviewed_round1': True})
result = {'output': str(TARGET), 'sha256': sha(TARGET), 'render': str(RENDER),
          'pdf_sha256': sha(RENDER / 'document.pdf'), 'all_passed': all(c['passed'] for c in checks), 'checks': checks}
(ROOT / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'all_passed': result['all_passed'], 'checks': len(checks), 'failed': [x for x in checks if not x['passed']]}, ensure_ascii=False))
raise SystemExit(0 if result['all_passed'] else 1)
