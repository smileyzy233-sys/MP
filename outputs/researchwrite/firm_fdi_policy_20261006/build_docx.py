from pathlib import Path
import re, copy, json, hashlib, os
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
PROJECT=ROOT.parents[2]
MS='http://www.w3.org/1998/Math/MathML'
NS={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
XSL=etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))

def node(tag,*children,text=None,**attrs):
    e=etree.Element('{'+MS+'}'+tag,**attrs)
    if text is not None: e.text=text
    for c in children:
        if c is not None: e.append(c)
    return e

GREEK={'alpha':'α','beta':'β','gamma':'γ','delta':'δ','epsilon':'ε','varepsilon':'ε','eta':'η','theta':'θ','kappa':'κ','lambda':'λ','mu':'μ','nu':'ν','xi':'ξ','pi':'π','rho':'ρ','sigma':'σ','tau':'τ','phi':'φ','psi':'ψ','omega':'ω','Delta':'Δ','Omega':'Ω','Psi':'Ψ'}
OPS={'in':'∈','notin':'∉','ne':'≠','neq':'≠','le':'≤','ge':'≥','equiv':'≡','simeq':'≈','approx':'≈','to':'→','times':'×','cdot':'·','mid':'|','partial':'∂','forall':'∀','prime':'′'}

class TexParser:
    def __init__(self,s,display=False):
        self.s=s
        self.tokens=re.findall(r'\\\\|\\[A-Za-z]+|\\.|[A-Za-z]|[0-9]+(?:\.[0-9]+)?|[^\s]',s)
        self.i=0; self.display=display
    def peek(self): return self.tokens[self.i] if self.i<len(self.tokens) else None
    def take(self):
        t=self.peek()
        if t is None: raise ValueError('Unexpected end: '+self.s)
        self.i+=1; return t
    def seq(self,stops=None):
        children=[]; stops=stops or set()
        while self.peek() is not None and self.peek() not in stops:
            a=self.atom()
            if a is None: continue
            sub=sup=None
            while self.peek() in ['_','^',"'"]:
                k=self.take()
                if k=="'": val=node('mo',text='′'); sup=val if sup is None else node('mrow',sup,val)
                else:
                    val=self.atom()
                    if k=='_': sub=val
                    else: sup=val
            if sub is not None or sup is not None:
                large=etree.QName(a).localname=='mo' and a.text in ['∑','∏'] and self.display
                if sub is not None and sup is not None: a=node('munderover' if large else 'msubsup',a,sub,sup)
                elif sub is not None: a=node('munder' if large else 'msub',a,sub)
                else: a=node('mover' if large else 'msup',a,sup)
            children.append(a)
        if len(children)==1:return children[0]
        return node('mrow',*children)
    def raw_group(self):
        if self.take()!='{': raise ValueError('Expected raw group')
        depth=1; items=[]
        while depth:
            t=self.take()
            if t=='{': depth+=1
            elif t=='}': depth-=1
            if depth: items.append(t)
        return ''.join(items)
    def atom(self):
        t=self.take()
        if t=='{':
            v=self.seq({'}'})
            if self.take()!='}': raise ValueError('Unclosed group')
            return v
        if t in ['}', '_','^']: raise ValueError('Unexpected '+t+' in '+self.s)
        if t.startswith('\\'):
            cmd=t[1:]
            if cmd in GREEK:return node('mi',text=GREEK[cmd])
            if cmd in OPS:return node('mo',text=OPS[cmd])
            if cmd in ['frac','dfrac','tfrac']:return node('mfrac',self.atom(),self.atom())
            if cmd in ['sum','prod']:return node('mo',text='∑' if cmd=='sum' else '∏',largeop='true',movablelimits='true')
            if cmd in ['exp','ln','log','min','max','arg']:
                return node('mi',text=cmd,mathvariant='normal')
            if cmd in ['hat','widehat','bar','overline']:
                return node('mover',self.atom(),node('mo',text='^' if cmd in ['hat','widehat'] else '¯'),accent='true')
            if cmd in ['mathbf','boldsymbol','mathcal','mathbb','mathrm','operatorname','text']:
                if cmd in ['text','operatorname']:
                    return node('mtext' if cmd=='text' else 'mi',text=self.raw_group(),mathvariant='normal')
                variant={'mathbf':'bold','boldsymbol':'bold-italic','mathcal':'script','mathbb':'double-struck','mathrm':'normal'}[cmd]
                return node('mstyle',self.atom(),mathvariant=variant)
            if cmd=='underset':
                lower=self.atom(); body=self.atom(); return node('munder',body,lower)
            if cmd in ['left','right']:
                d=self.take()
                if d=='.':return None
                if d in [r'\{',r'\}']:d=d[-1]
                return node('mo',text=d,stretchy='true')
            if cmd in [',',';','!',':',' ','quad','qquad']:
                return node('mspace',width='1em' if cmd=='quad' else '2em' if cmd=='qquad' else '0.16em')
            if cmd in ['{','}','_','%']:return node('mo',text=cmd)
            if cmd=='begin':
                env=self.raw_group()
                if env!='cases':raise ValueError('Unsupported environment '+env)
                rows=[]
                while self.peek()!=r'\end':
                    cells=[]
                    while True:
                        cells.append(node('mtd',self.seq({'&',r'\\',r'\end'})))
                        if self.peek()=='&':self.take();continue
                        break
                    rows.append(node('mtr',*cells))
                    if self.peek()==r'\\':self.take()
                    elif self.peek()!=r'\end':raise ValueError('Bad cases separator')
                self.take()
                if self.raw_group()!=env:raise ValueError('Environment mismatch')
                return node('mfenced',node('mtable',*rows,columnalign='left left',columnspacing='0.8em'),open='{',close='')
            raise ValueError('Unsupported command '+t+' in '+self.s)
        if re.fullmatch(r'\d+(\.\d+)?',t):return node('mn',text=t)
        if t.isalpha():return node('mi',text=t)
        return node('mo',text=t)
    def parse(self):
        result=self.seq()
        if self.i!=len(self.tokens):raise ValueError('Not fully parsed')
        return node('math',result,display='block' if self.display else 'inline')

math_records=[]
def math(tex,display=False):
    ml=TexParser(tex,display).parse()
    converted=XSL(ml).getroot()
    om=converted if etree.QName(converted).localname=='oMath' else converted.find('.//m:oMath',NS)
    if om is None:raise ValueError('No native equation')
    om=copy.deepcopy(om)
    for r in om.xpath('.//m:r',namespaces=NS):
        pr=r.find('w:rPr',NS)
        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
        fonts=OxmlElement('w:rFonts')
        for k,v in [('ascii','Cambria Math'),('hAnsi','Cambria Math'),('eastAsia','宋体')]:fonts.set(qn('w:'+k),v)
        pr.append(fonts)
        size=OxmlElement('w:sz');size.set(qn('w:val'),'23' if display else '24');pr.append(size)
    math_records.append({'latex':tex,'display':display,'mathml':etree.tostring(ml,encoding='unicode'),'native_text':''.join(om.xpath('.//m:t/text()',namespaces=NS))})
    return om

def font(style,size,east='宋体',bold=False):
    style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    pr=style.element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
    for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',east)]:rf.set(qn('w:'+k),v)
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+k),None)

def inline(p,text):
    for part in re.split(r'(\$[^$]+\$|`[^`]+`|\*\*[^*]+\*\*)',text):
        if not part:continue
        if part.startswith('$') and part.endswith('$'):p._p.append(math(part[1:-1]));continue
        code=part.startswith('`') and part.endswith('`')
        bold=part.startswith('**') and part.endswith('**')
        run=p.add_run(part[1:-1] if code else part[2:-2] if bold else part)
        if code:run.font.name='Consolas';run.font.size=Pt(10.5)
        if bold:run.bold=True

def add_table(doc,rows):
    n=len(rows[0]);tab=doc.add_table(rows=0,cols=n);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    ratios=[.19,.38,.43] if n==3 else [.4,.6] if n==2 else [1/n]*n
    full=Inches(6.5)
    for c,r in zip(tab.columns,ratios):c.width=int(full*r)
    bp=OxmlElement('w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        v=OxmlElement('w:'+side);v.set(qn('w:val'),'single' if side in ['top','bottom'] else 'nil');v.set(qn('w:sz'),'8');v.set(qn('w:color'),'000000');bp.append(v)
    tab._tbl.tblPr.append(bp)
    cm=OxmlElement('w:tblCellMar')
    for side,val in [('top',85),('bottom',85),('left',80),('right',80)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');cm.append(e)
    tab._tbl.tblPr.append(cm)
    for ri,row in enumerate(rows):
        rr=tab.add_row();pr=rr._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,txt) in enumerate(zip(rr.cells,row)):
            cell.width=int(full*ratios[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.style=doc.styles['Table Text'];inline(p,txt)
            if len(rows)<=4 and ri<len(rows)-1:p.paragraph_format.keep_with_next=True
            if ri==0:
                p.paragraph_format.keep_with_next=True
                for run in p.runs:run.bold=True
                cb=OxmlElement('w:tcBorders');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'000000');cb.append(b);cell._tc.get_or_add_tcPr().append(cb)
            if n==3 and ci==0 and len(txt)<=3:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=0.35

def build():
    draft=ROOT/'drafts/proposal_final.md';text=draft.read_text(encoding='utf-8')
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    for k in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,k,Inches(1))
    sec.footer_distance=Inches(.45)
    st=doc.styles;font(st['Normal'],12)
    st['Normal'].paragraph_format.line_spacing=1.4;st['Normal'].paragraph_format.space_after=Pt(5);st['Normal'].paragraph_format.widow_control=True
    font(st['Title'],18,'黑体',False);st['Title'].paragraph_format.space_after=Pt(9);st['Title'].paragraph_format.keep_with_next=True
    for name,size in [('Heading 1',14),('Heading 2',12.5)]:
        font(st[name],size,'黑体',False);p=st[name].paragraph_format;p.space_before=Pt(11 if name=='Heading 1' else 7);p.space_after=Pt(5);p.keep_with_next=True
    for name,size in [('Table Text',10.5),('Equation',11.5),('Reference',10.5),('Date Line',10.5)]:
        if name not in st:st.add_style(name,1)
        font(st[name],size)
    st['Table Text'].paragraph_format.line_spacing=1.15;st['Table Text'].paragraph_format.space_after=Pt(1)
    st['Equation'].paragraph_format.line_spacing=1.1;st['Equation'].paragraph_format.space_before=Pt(4);st['Equation'].paragraph_format.space_after=Pt(7);st['Equation'].paragraph_format.keep_together=True
    st['Equation'].paragraph_format.tab_stops.add_tab_stop(Inches(3.1),WD_TAB_ALIGNMENT.CENTER)
    st['Equation'].paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    st['Reference'].paragraph_format.line_spacing=1.2;st['Reference'].paragraph_format.space_after=Pt(5)
    font(st['Caption'],10.5);st['Caption'].paragraph_format.keep_with_next=True;st['Caption'].paragraph_format.space_before=Pt(6);st['Caption'].paragraph_format.space_after=Pt(4)
    for name in ['Title','Heading 1','Heading 2']:
        pr=st[name].element.get_or_add_pPr()
        for b in list(pr.findall(qn('w:pBdr'))):pr.remove(b)
    doc.core_properties.title='外资准入改革的企业效应与跨国生产成本反演研究方案'
    doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''
    mf=OxmlElement('m:mathPr');ff=OxmlElement('m:mathFont');ff.set(qn('m:val'),'Cambria Math');mf.append(ff);doc.settings.element.append(mf)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();run.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='1';rr.append(tt);fld.append(rr);footer._p.append(fld)
    lines=text.splitlines();i=0;refs=False;labels=[];captions=[];tables=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;inline(p,line[2:]);continue
        if line.startswith('## '):
            refs=line[3:]=='参考资料';doc.add_heading(line[3:],1);continue
        if line.startswith('### '):doc.add_heading(line[4:],2);continue
        if line=='2026年10月6日':
            p=doc.add_paragraph(style='Date Line');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(line);continue
        if line.startswith('$$'):
            expr=line[2:]
            while '$$' not in expr:
                if i>=len(lines):raise ValueError('Unclosed display')
                expr+='\n'+lines[i];i+=1
            expr=expr.rsplit('$$',1)[0].strip()
            label=re.search(r'\\tag\{([^}]+)\}',expr).group(1);expr=re.sub(r'\\tag\{[^}]+\}','',expr).strip()
            p=doc.add_paragraph(style='Equation');p.add_run('\t');p._p.append(math(expr,True));p.add_run('\t('+label+')');labels.append(label);continue
        if line.startswith('|'):
            rows=[]
            while True:
                vals=[v.strip() for v in line.strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v or '') for v in vals):rows.append(vals)
                if i>=len(lines) or not lines[i].strip().startswith('|'):break
                line=lines[i].strip();i+=1
            add_table(doc,rows);tables+=1;continue
        if re.match(r'^表A\d+ ',line):
            p=doc.add_paragraph(style='Caption');inline(p,line);captions.append(line);continue
        p=doc.add_paragraph(style='Reference' if refs else 'Normal');inline(p,line)
    assert labels==[str(x) for x in range(1,len(labels)+1)],labels
    meta_path=ROOT/'build_meta.json'
    if meta_path.exists():
        old=json.loads(meta_path.read_text(encoding='utf-8'));target=Path(old['output_path'])
        if target.exists() and hashlib.sha256(target.read_bytes()).hexdigest()!=old['output_sha256']:raise RuntimeError('Task output changed externally')
    else:
        target=PROJECT/'企业微观层面_新方案.docx';v=2
        while target.exists():target=PROJECT/f'企业微观层面_新方案_v{v}.docx';v+=1
    temp=ROOT/'exports'/'proposal_final.docx';temp.parent.mkdir(exist_ok=True);doc.save(temp)
    with ZipFile(temp) as z:
        assert z.testzip() is None
        xml=etree.fromstring(z.read('word/document.xml'));assert len(xml.xpath('//m:oMath',namespaces=NS))==len(math_records)
        assert not xml.xpath('//w:drawing',namespaces=NS),'Equations must remain native'
    os.replace(temp,target)
    md_target=target.with_suffix('.md')
    if md_target.exists() and not meta_path.exists():raise RuntimeError('Companion Markdown already exists')
    md_target.write_text(text,encoding='utf-8',newline='\n')
    meta={'output_path':str(target),'markdown_path':str(md_target),'output_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'draft_sha256':hashlib.sha256(draft.read_bytes()).hexdigest(),'native_math_objects':len(math_records),'display_equations':labels,'tables':tables,'captions':captions,'file_bytes':target.stat().st_size,'runtime':'Codex primary runtime 26.921.10847 bundled Python','math_conversion':'Task-local TeX subset to MathML and Microsoft Office MML2OMML.XSL'}
    meta_path.write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'equations.json').write_text(json.dumps(math_records,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(meta,ensure_ascii=False,indent=2))

if __name__=='__main__':build()
