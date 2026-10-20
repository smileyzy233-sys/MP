from pathlib import Path
from zipfile import ZipFile
import json, hashlib, re
from lxml import etree
from docx import Document
import pdfplumber

ROOT=Path(__file__).resolve().parent
meta=json.loads((ROOT/'build_meta.json').read_text(encoding='utf-8'))
target=Path(meta['output_path']);md=Path(meta['markdown_path'])
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
doc=Document(target)
with ZipFile(target) as z:
    assert z.testzip() is None
    root=etree.fromstring(z.read('word/document.xml'))
    assert len(root.xpath('//m:oMath',namespaces=ns))==meta['native_math_objects']
    prose=''.join(root.xpath('//w:t/text()',namespaces=ns))
    assert '$' not in prose
    assert not re.search(r'\\(?:frac|hat|sum|tag|begin)',prose)
    assert '\ufffd' not in prose
    images=len(root.xpath('//w:drawing',namespaces=ns))
    assert images==0
    mathtext=''.join(root.xpath('//m:t/text()',namespaces=ns))
    mathchars=root.xpath('//m:chr/@m:val',namespaces=ns)
    assert 'ξ' in mathtext and 'δ' in mathtext and '∑' in mathchars
heads=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
assert len(heads)==10
assert len(doc.tables)==5
assert md.read_bytes()==(ROOT/'drafts/proposal_final.md').read_bytes()
assert hashlib.sha256(target.read_bytes()).hexdigest()==meta['output_sha256'],'Render must be read-only'
pages=[]
with pdfplumber.open(ROOT/'qa_render/proposal_final.pdf') as pdf:
    for i,page in enumerate(pdf.pages,1):
        chars=[c for c in page.chars if c['text'].strip()]
        text=page.extract_text() or ''
        overflow=[c for c in chars if c['x0']<40 or c['x1']>page.width-35 or c['top']<35 or c['bottom']>page.height-25]
        pages.append({'page':i,'chars':len(text),'min_x':round(min(c['x0'] for c in chars),2),'max_x':round(max(c['x1'] for c in chars),2),'min_y':round(min(c['top'] for c in chars),2),'max_y':round(max(c['bottom'] for c in chars),2),'overflow_chars':len(overflow),'text':text,'fonts':sorted(set(c['fontname'] for c in chars))})
        assert not overflow,(i,overflow[:5])
        assert '\ufffd' not in text
        assert chars
    render_meta=json.loads((ROOT/'render_meta.json').read_text(encoding='utf-8-sig'))
    assert len(pdf.pages)==render_meta['pages']
source_hashes={}
expected={r'A:\桌面\study\企业微观层面.docx':'ec745fda88cac7840efc8fb5364e067c465da9125c577ee1ade9bcdec1949971',r'A:\桌面\制度性开放\审稿意见\手动改和写\审稿人回复.docx':'e82150a52e56d1cd276e669d5c300b2c55861b1d50202ddd35d0183529ffb608',r'A:\桌面\制度性开放\报告_重点标红版.pdf':'f6894e72569de840be2eadb4ea56a771207b0bde2b43758753c5b2501275d9d8'}
for p,h in expected.items():
    found=hashlib.sha256(Path(p).read_bytes()).hexdigest();assert found==h,(p,found);source_hashes[p]=found
pngs=sorted((ROOT/'qa_render').glob('page-*.png'));assert len(pngs)==len(pages)
summary={'docx_path':str(target),'bytes':target.stat().st_size,'pages':len(pages),'native_math':meta['native_math_objects'],'display_equations':14,'tables':5,'images':images,'source_hashes_unchanged':True,'sources':source_hashes,'docx_sha256':meta['output_sha256'],'markdown_matches_draft':True,'pdf_text_and_bounds_pass':True,'visual_review_pending':True}
(ROOT/'qa_render/page_metrics.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/'qa_logs').mkdir(exist_ok=True)
(ROOT/'qa_logs/verification.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))
for p in pages:print('PAGE',p['page'],'CHARS',p['chars'],'X',p['min_x'],p['max_x'],'Y',p['min_y'],p['max_y'])
