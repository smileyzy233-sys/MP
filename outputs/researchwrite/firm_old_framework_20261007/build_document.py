from pathlib import Path
from zipfile import ZipFile
import importlib.util,copy,re,json,hashlib
from lxml import etree
from docx import Document
from docx.shared import Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
SOURCE=Path(r'A:\桌面\study\企业微观层面_新方案.docx')
PROJECT=SOURCE.parent
sp=importlib.util.spec_from_file_location('native_math_helper',PROJECT/'outputs/researchwrite/firm_fdi_policy_20261006/build_docx.py')
helper=importlib.util.module_from_spec(sp);sp.loader.exec_module(helper)
NS=helper.NS;MS=helper.MS
records=[]

def native_math(tex,display=False):
    ml=helper.TexParser(tex,display).parse()
    # Office's converter reads leaf mathvariant; inherit styles explicitly.
    def inherit(e,current=None):
        local=e.get('mathvariant',current)
        if etree.QName(e).localname in ['mi','mo','mn','mtext'] and local:
            e.set('mathvariant',local)
        for child in e:inherit(child,local)
    inherit(ml)
    converted=helper.XSL(ml).getroot()
    om=converted if etree.QName(converted).localname=='oMath' else converted.find('.//m:oMath',NS)
    assert om is not None,tex
    om=copy.deepcopy(om)
    for r in om.xpath('.//m:r',namespaces=NS):
        pr=r.find('w:rPr',NS)
        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
        fonts=OxmlElement('w:rFonts')
        for k,v in [('ascii','Cambria Math'),('hAnsi','Cambria Math'),('eastAsia','宋体')]:fonts.set(qn('w:'+k),v)
        pr.append(fonts)
        size=OxmlElement('w:sz');size.set(qn('w:val'),'23' if display else '24');pr.append(size)
    records.append({'latex':tex,'display':display,'text':''.join(om.xpath('.//m:t/text()',namespaces=NS))})
    return om

def inline(p,text):
    for x in re.split(r'(\$[^$]+\$|`[^`]+`)',text):
        if not x:continue
        if x.startswith('$') and x.endswith('$'):p._p.append(native_math(x[1:-1]));continue
        code=x.startswith('`') and x.endswith('`')
        run=p.add_run(x[1:-1] if code else x)
        if code:run.font.name='Consolas';run.font.size=Pt(10.5)

def build():
    for e in json.loads((ROOT/'source_manifest.json').read_text(encoding='utf-8')):
        assert hashlib.sha256(Path(e['source']).read_bytes()).hexdigest()==e['sha256'],'Source changed'
    draft=(ROOT/'proposal_final.md').read_text(encoding='utf-8')
    doc=Document(SOURCE)
    body=doc.element.body
    for child in list(body):
        if etree.QName(child).localname!='sectPr':body.remove(child)
    # Retain the current document's page geometry, styles and page footer.
    doc.styles['Equation'].paragraph_format.tab_stops.clear_all()
    doc.styles['Equation'].paragraph_format.tab_stops.add_tab_stop(Inches(3.05),WD_TAB_ALIGNMENT.CENTER)
    doc.styles['Equation'].paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    doc.core_properties.title='附录A 外资准入负面清单的政策测度、跨国投资成本识别与一般均衡反事实'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    lines=draft.splitlines();i=0;refs=False;labels=[];heads=[];captions=[];tables=0
    # Reuse the proven table component, pointing its inline math to this writer.
    helper.inline=inline
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;inline(p,line[2:]);continue
        if line.startswith('## '):
            h=line[3:];refs=h=='参考文献';doc.add_heading(h,1);heads.append(h);continue
        if line.startswith('### '):doc.add_heading(line[4:],2);heads.append(line[4:]);continue
        if line.startswith('$$'):
            expr=line[2:]
            while '$$' not in expr:expr+='\n'+lines[i];i+=1
            expr=expr.rsplit('$$',1)[0].strip()
            tag=re.search(r'\\tag\{([^}]+)\}',expr);assert tag
            label=tag.group(1);expr=re.sub(r'\\tag\{[^}]+\}','',expr).strip()
            p=doc.add_paragraph(style='Equation');p.add_run('\t');p._p.append(native_math(expr,True));p.add_run('\t('+label+')');labels.append(label);continue
        if line.startswith('|'):
            rows=[]
            while True:
                vals=[s.strip() for s in line.strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',s or '') for s in vals):rows.append(vals)
                if i>=len(lines) or not lines[i].strip().startswith('|'):break
                line=lines[i].strip();i+=1
            helper.add_table(doc,rows);tables+=1;continue
        if re.match(r'^表A\d+',line):
            p=doc.add_paragraph(style='Caption');inline(p,line);p.paragraph_format.keep_with_next=True;captions.append(line);continue
        p=doc.add_paragraph(style='Reference' if refs else 'Normal');inline(p,line)
    assert labels==[f'A{n}' for n in range(1,18)],labels
    assert tables==3
    if (ROOT/'build_meta.json').exists():
        old=json.loads((ROOT/'build_meta.json').read_text(encoding='utf-8'));target=Path(old['output'])
        assert hashlib.sha256(target.read_bytes()).hexdigest()==old['sha256'],'Task output changed externally'
    else:
        target=PROJECT/'企业微观层面_新方案_框架重写版.docx';n=2
        while target.exists():target=PROJECT/f'企业微观层面_新方案_框架重写版_v{n}.docx';n+=1
    temp=ROOT/'pending.docx';doc.save(temp)
    with ZipFile(temp) as z:
        assert z.testzip() is None
        root=etree.fromstring(z.read('word/document.xml'))
        assert len(root.xpath('//m:oMath',namespaces=NS))==len(records)
        assert len(root.xpath('//m:sty[@m:val="b"]',namespaces=NS))>0,'Vector bold style missing'
        assert len(root.xpath('//m:scr[@m:val="script"]',namespaces=NS))>0,'Set notation missing'
    temp.replace(target)
    meta={'output':str(target),'sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'bytes':target.stat().st_size,'draft_sha256':hashlib.sha256((ROOT/'proposal_final.md').read_bytes()).hexdigest(),'main_sections':7,'subsections':16,'equation_labels':labels,'native_math_objects':len(records),'tables':tables,'table_captions':captions,'headings':heads,'style_reference':str(SOURCE),'framework_reference':r'A:\桌面\study\旧版方案\企业微观层面.docx'}
    (ROOT/'build_meta.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'equations.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(meta,ensure_ascii=False,indent=2))

if __name__=='__main__':build()
