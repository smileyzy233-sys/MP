from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
import hashlib,json,re
import pdfplumber

ROOT=Path(__file__).resolve().parent
meta=json.loads((ROOT/'build_meta.json').read_text(encoding='utf-8'))
target=Path(meta['output'])
for src in json.loads((ROOT/'source_manifest.json').read_text(encoding='utf-8')):
    assert hashlib.sha256(Path(src['source']).read_bytes()).hexdigest()==src['sha256'],'Source modified'
doc=Document(target)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
with ZipFile(target) as z:
    assert z.testzip() is None
    root=etree.fromstring(z.read('word/document.xml'))
    assert len(root.xpath('//m:oMath',namespaces=ns))==meta['native_math_objects']
    assert not root.xpath('//w:drawing',namespaces=ns),'Math must be native'
    prose=''.join(root.xpath('//w:t/text()',namespaces=ns))
    assert '$' not in prose and '\ufffd' not in prose
    assert not re.search(r'\\(?:frac|hat|sum|mathcal|boldsymbol)',prose)
    assert 'X_{nCmj}' not in prose
    assert 'IMPC=' not in prose and '历史制度成本存量' not in prose
    assert root.xpath('//m:rPr/m:sty[@m:val="b"]',namespaces=ns)
    assert root.xpath('//m:rPr/m:scr[@m:val="script"]',namespaces=ns)
assert len(doc.tables)==3
main=[p.text for p in doc.paragraphs if p.style.name=='Heading 1' and p.text.startswith('A.')]
subs=[p.text for p in doc.paragraphs if p.style.name=='Heading 2']
assert len(main)==7 and len(subs)==16
assert [re.match(r'A\.(\d+)',s).group(1) for s in main]==[str(x) for x in range(1,8)]
labels=[p.text.strip() for p in doc.paragraphs if p.style.name=='Equation']
assert [re.search(r'\((A\d+)\)',s).group(1) for s in labels]==[f'A{x}' for x in range(1,18)]
pages=[]
with pdfplumber.open(ROOT/'render/document.pdf') as pdf:
    for i,p in enumerate(pdf.pages,1):
        chars=[c for c in p.chars if c['text'].strip()]
        t=p.extract_text() or ''
        bad=[c for c in chars if c['x0']<35 or c['x1']>p.width-25 or c['top']<25 or c['bottom']>p.height-15]
        assert chars and not bad,(i,bad[:5])
        assert '\ufffd' not in t
        pages.append({'page':i,'chars':len(t),'min_x':min(c['x0'] for c in chars),'max_x':max(c['x1'] for c in chars),'min_y':min(c['top'] for c in chars),'max_y':max(c['bottom'] for c in chars),'text':t})
render=json.loads((ROOT/'render/render_meta.json').read_text(encoding='utf-8-sig'))
assert len(pages)==render['pages']
assert len(list((ROOT/'render').glob('page-*.png')))==len(pages)
assert hashlib.sha256(target.read_bytes()).hexdigest()==meta['sha256'],'Read-only render changed output'
result={'output':str(target),'sha256':meta['sha256'],'pages':len(pages),'main_sections':7,'subsections':16,'tables':3,'display_formulas':17,'native_math':meta['native_math_objects'],'source_files_unchanged':True,'native_math_styles_passed':True,'word_opens':True,'pdf_bounds_passed':True,'visual_review_pending':True}
(ROOT/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/'page_metrics.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
print('PAGE_CHARS',[(x['page'],x['chars']) for x in pages])
