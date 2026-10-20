from pathlib import Path
import re, json, hashlib, copy
from datetime import datetime
from lxml import etree
import latex2mathml.converter
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[3]
BASE = Path(__file__).resolve().parent
SOURCE = ROOT / '企业微观层面_v3.3.docx'
EXPECTED = '6caf8e57849bc454f5a695c412ab962680e5460cf1fa310cb0df61d3eefc017c'
XSL = etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
NS = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math', 'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
MML = {'x': 'http://www.w3.org/1998/Math/MathML'}
PAGE_BREAK_BEFORE = {'### A.7.2 模型基期与生产口径', '### A.7.4 反事实设置、边界与实施顺序'}
MATH_LOG = []

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def no_keeps(p):
    pf = p.paragraph_format
    pf.keep_with_next = False
    pf.keep_together = False
    pf.page_break_before = False

def font(run, size=12, bold=False, face='Times New Roman', east='宋体', color='000000'):
    run.font.name = face
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    for a, v in [('ascii', face), ('hAnsi', face), ('eastAsia', east), ('cs', face)]:
        fonts.set(qn('w:' + a), v)
    return run

def math_el(tex, size=12):
    tex = tex.replace('Treat', r'\mathrm{Treat}').replace('Post', r'\mathrm{Post}').replace('MP', r'\mathrm{MP}')
    tex = tex.replace(r'\qquad', r'\text{  }').replace(r'\quad', r'\text{ }')
    xml = etree.fromstring(latex2mathml.converter.convert(tex).encode('utf-8'))
    for mover in xml.xpath('.//x:mover', namespaces=MML):
        if len(mover) == 2 and ''.join(mover[1].itertext()) in ['^', 'ˆ', '¯', '→']:
            mover.set('accent', 'true')
    om = etree.fromstring(etree.tostring(XSL(xml)))
    unexplained = [t.text for t in om.xpath('.//m:t', namespaces=NS) if t.text and '\\' in t.text]
    if unexplained:
        raise ValueError(f'Unconverted TeX {tex}: {unexplained}')
    for t in om.xpath('.//m:t', namespaces=NS):
        if t.text and t.text.isspace():t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    for r in om.xpath('.//m:r', namespaces=NS):
        if ''.join(r.xpath('./m:t/text()', namespaces=NS)) == 'argmin':
            mrp = r.find(qn('m:rPr'))
            if mrp is None:mrp=OxmlElement('m:rPr');r.insert(0,mrp)
            sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');mrp.append(sty)
        rp = OxmlElement('w:rPr')
        f = OxmlElement('w:rFonts')
        for attr in ['ascii', 'hAnsi', 'cs']: f.set(qn('w:' + attr), 'Cambria Math')
        f.set(qn('w:eastAsia'), '宋体'); rp.append(f)
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rp.append(sz)
        # Math run properties precede Word character properties.
        pos = 1 if len(r) and r[0].tag == qn('m:rPr') else 0
        r.insert(pos, rp)
    MATH_LOG.append({'tex': tex, 'text': ''.join(om.xpath('.//m:t/text()', namespaces=NS))})
    return om

TOKEN = re.compile(r'(\$[^$]+\$|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))')

def inline(p, text, size=12, bold=False):
    for token in TOKEN.split(text):
        if not token: continue
        if token.startswith('$') and token.endswith('$'):
            p._p.append(math_el(token[1:-1], size))
        elif token.startswith('`') and token.endswith('`'):
            font(p.add_run(token[1:-1]), size-0.7, bold, 'Consolas')
        elif re.fullmatch(r'\[[^\]]+\]\(https?://[^)]+\)', token):
            label, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token).groups()
            h = OxmlElement('w:hyperlink')
            h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = p.add_run(label); font(run, size, bold, color='234A70')
            h.append(run._r); p._p.append(h)
        else:
            font(p.add_run(token), size, bold)

def ordinary_break(doc):
    children=[e for e in doc._element.body if e.tag!=qn('w:sectPr')]
    if children and children[-1].tag==qn('w:p'):
        prior=Paragraph(children[-1],doc._body)
        font(prior.add_run(),1).add_break(WD_BREAK.PAGE)
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    p.paragraph_format.widow_control = False
    p.paragraph_format.first_line_indent = Pt(0)
    font(p.add_run(), 1).add_break(WD_BREAK.PAGE)
    no_keeps(p)

def setup(note=False):
    d = Document()
    s = d.sections[0]
    s.page_height = Cm(29.7); s.page_width = Cm(21)
    s.top_margin = Cm(2.15); s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.35); s.right_margin = Cm(2.35)
    s.header_distance = Cm(.9); s.footer_distance = Cm(.9)
    normal = d.styles['Normal']
    normal.font.name = 'Times New Roman'; normal.font.size = Pt(11.5 if note else 12)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    pf = normal.paragraph_format
    pf.line_spacing = 1.3 if note else 1.35
    pf.space_after = Pt(5 if note else 4)
    pf.first_line_indent = Pt(23 if note else 24)
    pf.widow_control = True
    for nm, size in [('Title', 18), ('Subtitle', 10.5), ('Heading 1', 14), ('Heading 2', 12.5)]:
        st = d.styles[nm]
        st.font.name = 'Times New Roman'; st.font.size = Pt(size)
        st.font.bold = nm != 'Subtitle'; st.font.italic=False; st.font.color.rgb = RGBColor(0,0,0)
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体' if nm in ['Title','Heading 1'] else '宋体')
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.line_spacing = 1.2
        st.paragraph_format.space_before = Pt(12 if nm.startswith('Heading') else 0)
        st.paragraph_format.space_after = Pt(6)
    # The Word default heading styles use keepNext; explicitly disable for all styles.
    for st in d.styles:
        if hasattr(st, 'paragraph_format'):
            for attr in ['keep_with_next', 'keep_together', 'page_break_before']:
                setattr(st.paragraph_format, attr, False)
        for border in list(st._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.core_properties.title = '企业微观层面 v3.35 修订说明' if note else '2011年外资准入放宽的企业效应与部门成本识别'
    d.core_properties.subject = '2011年外资准入改革研究方案'
    d.core_properties.author = ''
    d.core_properties.last_modified_by = ''
    d.core_properties.keywords = '2011; 外资准入; DID; AMNE; 跨国生产'
    hp = s.header.paragraphs[0]
    hp.paragraph_format.first_line_indent = Pt(0)
    hp.paragraph_format.space_after = Pt(0)
    font(hp.add_run('企业微观层面 · v3.35' + (' · 修订说明' if note else '')),9,color='666666')
    no_keeps(hp)
    fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Pt(0)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    r=OxmlElement('w:r'); t=OxmlElement('w:t');t.text='1';r.append(t);fld.append(r);fp._p.append(fld)
    no_keeps(fp)
    return d

def set_table_border(cell, edge, value='single', sz=6, color='666666'):
    pr=cell._tc.get_or_add_tcPr(); borders=pr.find(qn('w:tcBorders'))
    if borders is None: borders=OxmlElement('w:tcBorders');pr.append(borders)
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),value);el.set(qn('w:sz'),str(sz));el.set(qn('w:color'),color);borders.append(el)

def table(doc, rows, widths=None):
    tab=doc.add_table(rows=len(rows),cols=len(rows[0])); tab.alignment=WD_TABLE_ALIGNMENT.CENTER
    tab.autofit=False
    if widths is None: widths=[16.3/len(rows[0])]*len(rows[0])
    for rowi,row in enumerate(tab.rows):
        trpr=row._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit');trpr.append(cs)
        if rowi==0: repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.line_spacing=1.18;no_keeps(p)
            inline(p,rows[rowi][j],10.5,rowi==0)
            for edge in ['left','right','top','bottom']:set_table_border(cell,edge,'nil')
            if rowi==0:
                set_table_border(cell,'top',sz=10,color='000000');set_table_border(cell,'bottom',sz=6,color='000000')
            if rowi==len(rows)-1:set_table_border(cell,'bottom',sz=10,color='000000')
    for col, width in zip(tab.columns,widths):col.width=Cm(width)
    return tab

def flowchart(doc):
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    inline(p,'图A1  政策识别、部门归属与研究层次',11,True)
    t=doc.add_table(rows=4,cols=2);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            c.width=Cm(8.15)
            for edge in ['left','right','top','bottom']:set_table_border(c,edge,'single',4,'999999')
    texts=[
        ['路径A：固定政策处理身份','路径B：确定当期生产归属'],
        ['2010年或更早的真实业务\n或已核实版本的行业定义\n↓\n原条款、前后要求及适用条件\n↓\n固定Treat或未知状态',
         '各期原始行业及分类版本\n↓\n必要的新旧转换 → GB2017\n↓\nISIC Rev.4 → AMNE → 模型部门'],
    ]
    for i in range(2):
        for j in range(2):
            c=t.cell(i,j);p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.2
            p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
            inline(p,texts[i][j],10.5,i==0)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EFEFEF');c._tc.get_or_add_tcPr().append(sh)
    for i,text in [(2,'↓  企业DID与生产统计：核对比较组、控制关系、生产口径和覆盖范围'),(3,'↓  2014年可比生产目标 → 等效MP成本 → 一般均衡反事实\n后续环节以资料和识别条件成立为前提')]:
        c=t.cell(i,0).merge(t.cell(i,1));p=c.paragraphs[0]
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6)
        inline(p,text,10.5)
    return t

def equation(doc, line):
    tex=line[2:-2]
    match=re.search(r'\\tag\{([^}]+)\}',tex)
    tag=match.group(1);tex=tex[:match.start()].strip()
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.line_spacing=1.05
    p.paragraph_format.tab_stops.add_tab_stop(Cm(7.6),WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.3),WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run('\t'))
    p._p.append(math_el(tex,11 if tag in ['A3','A8','A9'] else 12))
    font(p.add_run('\t（'+tag+'）'),11)
    no_keeps(p)

def finalize(doc):
    # Remove inherited/default enabling flags in all Word parts, then explicitly disable.
    for part in doc.part.package.parts:
        generic=not hasattr(part,'_element')
        if generic:
            if not str(part.partname).startswith('/word/') or not str(part.partname).endswith('.xml'):continue
            root=etree.fromstring(part.blob)
        else:root=part._element
        for tag in ['keepNext','keepLines','pageBreakBefore']:
            for e in root.iter(qn('w:'+tag)):e.set(qn('w:val'),'0')
        for p in root.iter(qn('w:p')):
            pp=p.find(qn('w:pPr'))
            if pp is None:pp=OxmlElement('w:pPr');p.insert(0,pp)
            for tag in ['keepNext','keepLines','pageBreakBefore']:
                e=pp.find(qn('w:'+tag))
                if e is None:e=OxmlElement('w:'+tag);pp.append(e)
                e.set(qn('w:val'),'0')
        if generic:part._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)

def build(md, output, note=False):
    doc=setup(note);lines=md.read_text(encoding='utf-8').splitlines();i=0;first_section=True
    while i<len(lines):
        s=lines[i].strip();i+=1
        if not s:continue
        if s=='[[FLOWCHART]]':flowchart(doc);continue
        if s.startswith('$$'):equation(doc,s);continue
        if s.startswith('|'):
            tbl=[s]
            while i<len(lines) and lines[i].strip().startswith('|'):tbl.append(lines[i].strip());i+=1
            rows=[[c.strip() for c in line.strip('|').split('|')] for line in tbl if not re.fullmatch(r'[| :\-]+',line)]
            widths=[2.5,4.1,9.7] if rows[0][0]=='年度' else [3.5,5.3,7.5]
            table(doc,rows,widths);continue
        if s.startswith('# '):
            p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,s[2:],18,True);p.paragraph_format.space_after=Pt(10)
        elif s.startswith('## '):
            if not note and (s.startswith('## A.1 ') or s=='## 参考文献与资料说明'):ordinary_break(doc)
            first_section=False
            p=doc.add_paragraph(style='Heading 1');inline(p,s[3:],14,True)
            p.paragraph_format.space_before=Pt(0 if not note else 9)
        elif s.startswith('### '):
            if s in PAGE_BREAK_BEFORE:ordinary_break(doc)
            p=doc.add_paragraph(style='Heading 2');inline(p,s[4:],12.5,True)
        elif i==2:
            p=doc.add_paragraph(style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,s,10.5);p.paragraph_format.space_after=Pt(14)
        else:
            if s.startswith('表A1 ') or s=='静态政策项为' or s.startswith(r'令$G_{\chi}'):
                ordinary_break(doc)
            p=doc.add_paragraph()
            if s.startswith('表A'):
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
                inline(p,s,11,True)
            elif s.startswith('表注：'):
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15
                p.paragraph_format.space_before=Pt(4);inline(p,s,10)
            elif re.match(r'^\[\d+\]',s):
                p.paragraph_format.first_line_indent=Pt(-18);p.paragraph_format.left_indent=Pt(18)
                p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(7)
                inline(p,s,10.5)
            else:inline(p,s,11.5 if note else 12)
        no_keeps(p)
    finalize(doc)
    doc.save(output)
    return {'path':str(output),'sha256':sha(output),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)}

if __name__=='__main__':
    assert sha(SOURCE)==EXPECTED,'Original source changed; review before building'
    targets=[ROOT/'企业微观层面_v3.35.docx',ROOT/'企业微观层面_v3.35_修订说明.docx']
    prior=json.loads((BASE/'build_meta.json').read_text(encoding='utf-8')) if (BASE/'build_meta.json').exists() else {}
    priorhash={item['path']:item['sha256'] for item in prior.get('outputs',[])}
    for target in targets:
        if target.exists():assert priorhash.get(str(target))==sha(target),'Existing target not from last build; preserve it'
    results=[build(BASE/'proposal_v335.md',targets[0]),build(BASE/'revision_note.md',targets[1],True)]
    assert sha(SOURCE)==EXPECTED
    meta={'source':str(SOURCE),'source_sha256':EXPECTED,'outputs':results,'math_objects':len(MATH_LOG),'draft_sha256':sha(BASE/'proposal_v335.md'),'generated_at':datetime.now().isoformat(),'math_engine':'latex2mathml + Microsoft MML2OMML; native OMML'}
    (BASE/'build_meta.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    (BASE/'qa'/'math_manifest.json').write_text(json.dumps(MATH_LOG,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(meta,ensure_ascii=False,indent=2))
